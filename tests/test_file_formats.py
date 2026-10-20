
import json
import csv

import pytest
from docx import Document
from openpyxl import Workbook

from src.preprocessing.pipeline import preprocess_document


@pytest.fixture
def sample_files(tmp_path):
    files = {}

    # 1. CSV
    csv_path = tmp_path / "claims.csv"
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["claim_id", "hospital", "amount"])
        writer.writerow(["CLM001", "Hospital A", "4500"])
        writer.writerow(["CLM002", "Hospital B", "3200"])
    files["csv"] = csv_path

    # 2. Excel
    excel_path = tmp_path / "claims.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Claims"
    sheet.append(["claim_id", "hospital", "amount"])
    sheet.append(["CLM001", "Hospital A", 4500])
    sheet.append(["CLM002", "Hospital B", 3200])
    workbook.save(excel_path)
    workbook.close()
    files["xlsx"] = excel_path

    # 3. JSON
    json_path = tmp_path / "claims.json"
    data = {
        "claim_id": "CLM001",
        "hospital": "Hospital A",
        "amount": 4500
    }
    json_path.write_text(json.dumps(data), encoding="utf-8")
    files["json"] = json_path

    # 4. TXT
    txt_path = tmp_path / "claim.txt"
    txt_path.write_text(
        "Claim ID: CLM001\n"
        "Hospital: Hospital A\n"
        "Amount: 4500",
        encoding="utf-8"
    )
    files["txt"] = txt_path

    # 5. Word
    docx_path = tmp_path / "claim.docx"
    document = Document()
    document.add_heading("Hospital Claim", level=1)
    document.add_paragraph("Claim ID: CLM001")
    document.add_paragraph("Hospital: Hospital A")
    document.add_paragraph("Amount: 4500")
    document.save(docx_path)
    files["docx"] = docx_path

    return files


@pytest.mark.parametrize(
    "file_type",
    ["csv", "xlsx", "json", "txt", "docx"]
)
def test_file_ingestion(sample_files, file_type, tmp_path):
    file_path = sample_files[file_type]
    output_path = tmp_path / f"{file_type}_processed.json"

    result = preprocess_document(
        file_path=file_path,
        output_path=output_path
    )

    # Check common document structure
    assert result["filename"] == file_path.name
    assert result["file_type"] == file_type
    assert result["document_id"]
    assert result["pages"]
    assert result["page_count"] == 1

    # Check text was extracted
    text = result["pages"][0]["text"]
    assert "CLM001" in text
    assert "Hospital A" in text

    # Check processing status based on file type
    if file_type in {"txt", "docx"}:
        assert result["processing_status"] == "processed"
        assert "classification" in result
        assert "extraction" in result

    else:
        assert result["processing_status"] == "ingested"
        assert result["chunk_count"] == 0
        assert result["extraction"]["field_count"] == 0

    # Check output JSON was saved
    assert output_path.exists()
    saved = json.loads(output_path.read_text(encoding="utf-8"))
    assert saved["document_id"] == result["document_id"]

    print(f"{file_type.upper()} ingestion passed")


def test_csv_structure(sample_files):
    result = preprocess_document(sample_files["csv"])
    table = result["structured_data"]

    assert table["row_count"] == 2
    assert table["columns"] == [
        "claim_id", "hospital", "amount"
    ]
    assert table["rows"][0]["values"]["claim_id"] == "CLM001"


def test_excel_structure(sample_files):
    result = preprocess_document(sample_files["xlsx"])
    data = result["structured_data"]

    assert data["sheet_count"] == 1
    assert data["sheets"][0]["sheet_name"] == "Claims"
    assert data["sheets"][0]["row_count"] == 3
def test_json_structure(sample_files):
    result = preprocess_document(sample_files["json"])
    data = result["structured_data"]

    assert data["claim_id"] == "CLM001"
    assert data["amount"] == 4500


def test_word_structure(sample_files):
    result = preprocess_document(sample_files["docx"])
    text = result["pages"][0]["text"]

    assert "Hospital Claim" in text
    assert "Claim ID: CLM001" in text


def test_unsupported_format(tmp_path):
    path = tmp_path / "invalid.xyz"
    path.write_text("Invalid file", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        preprocess_document(path)
