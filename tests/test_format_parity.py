"""No SLM is required for deterministic record classification tests."""
from docx import Document

from src.extraction.structured_record_classifier import (
    classify_structured_record,
)
from src.preprocessing.document_builder import process_document
from src.extraction.structured_field_mapper import make_record


def _record(values):
    locations = {
        key: {"type": "csv", "row": 2, "column": key}
        for key in values
    }
    return make_record("row_2", values, locations)


def test_shared_claim_fields_are_not_assumed_hospital_bill():
    result = classify_structured_record(_record({
        "claim_id": "CLM001",
        "patient_name": "Alex Morgan",
        "hospital_name": "General Hospital",
        "total_amount": 8700,
    }))
    assert result["document_type"] == "unknown"
    assert result["documentary_evidence"] is False
    assert result["requires_review"] is True


def test_distinctive_schema_is_heuristic_only():
    result = classify_structured_record(_record({
        "claim_id": "CLM001",
        "patient_name": "Alex Morgan",
        "prescription_number": "RX001",
        "medication_name": "medicine",
    }))
    assert result["document_type"] == "prescription"
    assert result["classification_method"] == "schema_heuristic"
    assert result["documentary_evidence"] is False
    assert result["requires_review"] is True


def test_conflicting_schema_markers_stay_unknown():
    result = classify_structured_record(_record({
        "bill_number": "B1",
        "prescription_number": "RX001",
    }))
    assert result["document_type"] == "unknown"


def test_docx_preserves_heading_style_and_table_location(tmp_path):
    path = tmp_path / "sample.docx"
    word = Document()
    word.add_heading("Treatment Summary", level=1)
    word.add_paragraph("Patient Name: Alex Morgan")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Medicine"
    table.cell(0, 1).text = "Amount"
    word.save(path)

    document = process_document(path)
    lines = document["pages"][0]["lines"]

    assert lines[0]["is_bold"] is True
    assert lines[0]["source"]["type"] == "docx_paragraph"
    assert any(line["source"]["type"] == "docx_table_row"
               for line in lines)
    assert lines[0]["text"] == "Treatment Summary"


def test_txt_preserves_line_numbers(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_text("First line\nSecond line\n", encoding="utf-8")
    document = process_document(path)
    lines = document["pages"][0]["lines"]
    assert lines[0]["source"]["line"] == 1
    assert lines[1]["source"]["line"] == 2
