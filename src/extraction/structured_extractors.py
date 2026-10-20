
import csv
import json
from datetime import date, datetime
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def to_json_value(value):
    """Convert spreadsheet values into JSON-compatible values."""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def extract_csv(file_path):
    with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)

        try:
            dialect = csv.Sniffer().sniff(
                sample, delimiters=",;\t"
            )
        except csv.Error:
            dialect = csv.excel

        reader = csv.DictReader(f, dialect=dialect)

        if not reader.fieldnames:
            raise ValueError("CSV is missing column headers")

        if any(not h or not h.strip() for h in reader.fieldnames):
            raise ValueError("CSV contains empty column headers")

        if len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError("CSV contains duplicate headers")

        rows = []
        for row_number, row in enumerate(reader, start=2):
            if None in row:
                raise ValueError(f"Extra columns in row {row_number}")

            rows.append({
                "row_number": row_number,
                "values": row
            })

    columns = reader.fieldnames

    lines = [
        (f"Row {row['row_number']}: " +
         " | ".join(f"{k}: {v}" for k, v in row["values"].items()))
        for row in rows
    ]

    return {
        "text": "\n".join(lines),
        "structured_data": {
            "columns": columns,
            "rows": rows,
            "row_count": len(rows)
        }
    }


def extract_excel(file_path):
    workbook = load_workbook(
        file_path, read_only=True, data_only=True
    )

    sheets = []
    text_lines = []

    try:
        for sheet in workbook.worksheets:
            sheet_rows = []

            for row_number, row in enumerate(
                sheet.iter_rows(values_only=True), start=1
            ):
                values = [to_json_value(v) for v in row]

                sheet_rows.append({
                    "row_number": row_number,
                    "values": values
                })

                text_lines.append(
                    f"Sheet {sheet.title}, Row {row_number}: " +
                    " | ".join(
                        "" if v is None else str(v)
                        for v in values
                    )
                )

            sheets.append({
                "sheet_name": sheet.title,
                "rows": sheet_rows,
                "row_count": len(sheet_rows)
            })
    finally:
        workbook.close()

    return {
        "text": "\n".join(text_lines),
        "structured_data": {
            "sheets": sheets,
            "sheet_count": len(sheets)
        }
    }


def extract_json(file_path):
    with open(file_path, "r", encoding="utf-8-sig") as f:
        data = json.load(f)

    return {
        "text": json.dumps(
            data, indent=2, ensure_ascii=False
        ),
        "structured_data": data
    }


def extract_txt(file_path):
    text = Path(file_path).read_text(encoding="utf-8-sig")

    return {
        "text": text,
        "structured_data": None
    }


def extract_word(file_path):
    document = Document(file_path)
    content = []
    line_metadata = []
    tables = []

    # Keep paragraph and table order as it appears in the document body.
    # A logical 'page' is not a physical Word page.
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    paragraph_num = 0
    table_num = 0

    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            paragraph_num += 1
            if not paragraph.text.strip():
                continue
            style = paragraph.style.name if paragraph.style else ""
            is_heading = style.lower().startswith(("heading", "title", "subtitle"))
            is_bold = is_heading or bool(
                paragraph.runs and all(
                    run.bold is True for run in paragraph.runs if run.text.strip()
                )
            )
            content.append(paragraph.text)
            line_metadata.append({
                "type": "docx_paragraph",
                "paragraph_number": paragraph_num,
                "style": style,
                "is_bold": is_bold,
            })
        elif element.tag == qn("w:tbl"):
            table_num += 1
            table = Table(element, document)
            table_rows = []
            content.append(f"Table {table_num}")
            line_metadata.append({"type": "docx_table", "table": table_num})

            for row_number, row in enumerate(table.rows, start=1):
                values = [cell.text for cell in row.cells]
                table_rows.append(values)
                content.append(
                    f"Row {row_number}: " + " | ".join(values)
                )
                line_metadata.append({
                    "type": "docx_table_row",
                    "table": table_num,
                    "row": row_number,
                })
            tables.append({
                "table_number": table_num,
                "rows": table_rows,
            })

    return {
        "text": "\\n".join(content),
        "structured_data": {"tables": tables},
        "line_metadata": line_metadata,
    }


EXTRACTORS = {
    ".csv": extract_csv,
    ".xlsx": extract_excel,
    ".json": extract_json,
    ".txt": extract_txt,
    ".docx": extract_word,
}


def extract_non_pdf(file_path):
    file_path = Path(file_path)
    extension = file_path.suffix.lower()

    if extension not in EXTRACTORS:
        raise ValueError(f"Unsupported file type: {extension}")

    return EXTRACTORS[extension](file_path)
